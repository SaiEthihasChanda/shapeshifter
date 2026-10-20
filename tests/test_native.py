"""Native (from-scratch) codecs, checked against Pillow as the reference."""

from __future__ import annotations

import io
import struct
import sys
import wave
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PIL import Image  # noqa: E402

from shapeshifter import config, engine, planner  # noqa: E402
from shapeshifter.backends import Context  # noqa: E402
from shapeshifter.native import audio, docs, gif, ico, jpeg, png, simple, xlsx  # noqa: E402
from shapeshifter.native.raster import Frame, Raster, Unsupported  # noqa: E402

CFG = dict(config.DEFAULTS, original="keep")


def gradient(mode="RGB", size=(37, 23)) -> Image.Image:
    w, h = size
    im = Image.new("RGBA", size)
    im.putdata([(x * 7 % 256, y * 11 % 256, (x * y) % 256, (x + y) * 5 % 256)
                for y in range(h) for x in range(w)])
    return im.convert(mode)


def pil_bytes(im: Image.Image, fmt: str, **kw) -> bytes:
    b = io.BytesIO()
    im.save(b, fmt, **kw)
    return b.getvalue()


def same_pixels(raster: Raster, im: Image.Image, tol: int = 0):
    f = raster.first
    ref = im.convert(f.mode).tobytes()
    assert (f.width, f.height) == im.size
    if tol == 0:
        assert f.data == ref
    else:
        diff = max(abs(a - b) for a, b in zip(f.data, ref))
        mean = sum(abs(a - b) for a, b in zip(f.data, ref)) / len(ref)
        assert mean < tol / 2 and diff <= tol * 2, (mean, diff)   # IDCT rounding only


# --------------------------------------------------------------------- decoders

@pytest.mark.parametrize("mode", ["RGB", "RGBA", "L", "LA", "P", "1"])
def test_png_decode_modes(mode):
    im = gradient("RGB").convert(mode) if mode != "LA" else gradient("LA")
    same_pixels(png.decode(pil_bytes(im, "PNG")), im.convert(
        {"P": "RGB", "1": "L"}.get(mode, mode)))


def test_png_decode_palette_transparency_and_16bit():
    im = gradient("RGBA").quantize(64)            # P with tRNS
    r = png.decode(pil_bytes(im, "PNG"))
    same_pixels(r, im.convert("RGBA"))
    im16 = gradient("L").convert("I;16")
    r = png.decode(pil_bytes(im16, "PNG"))
    assert r.first.mode == "L" and (r.first.width, r.first.height) == im16.size


@pytest.mark.parametrize("mode", ["RGB", "RGBA"])
def test_png_encode_round_trip(mode):
    im = gradient(mode)
    data = png.encode(Raster([Frame(*im.size, mode, im.tobytes())]))
    assert Image.open(io.BytesIO(data)).convert(mode).tobytes() == im.tobytes()


@pytest.mark.parametrize("fmt,dec,enc,mode", [
    ("BMP", simple.decode_bmp, simple.encode_bmp, "RGB"),
    ("BMP", simple.decode_bmp, simple.encode_bmp, "RGBA"),
    ("TGA", simple.decode_tga, simple.encode_tga, "RGB"),
    ("TGA", simple.decode_tga, simple.encode_tga, "RGBA"),
    ("PPM", simple.decode_pnm, simple.encode_pnm, "RGB"),
    ("PPM", simple.decode_pnm, simple.encode_pnm, "L"),
    ("QOI", simple.decode_qoi, simple.encode_qoi, "RGB"),
    ("QOI", simple.decode_qoi, simple.encode_qoi, "RGBA"),
])
def test_simple_formats_against_pillow(fmt, dec, enc, mode):
    im = gradient(mode)
    if not (fmt == "BMP" and mode == "RGBA"):
        same_pixels(dec(pil_bytes(im, fmt)), im)          # we read what Pillow writes
    ours = enc(Raster([Frame(*im.size, mode, im.tobytes())]))
    back = Image.open(io.BytesIO(ours))
    assert back.size == im.size
    assert back.convert(mode).tobytes() == im.tobytes()   # Pillow reads what we write
    same_pixels(dec(ours), im)                            # and we read our own


def test_bmp_palette_and_rle_tga():
    im = gradient("RGB").convert("P", palette=Image.ADAPTIVE, colors=16)
    same_pixels(simple.decode_bmp(pil_bytes(im, "BMP")), im.convert("RGB"))
    rgb = gradient("RGB")
    same_pixels(simple.decode_tga(pil_bytes(rgb, "TGA", compression="tga_rle")), rgb)


def test_gif_animation_matches_pillow():
    frames = [gradient("RGB", (30, 20)).rotate(i * 40) for i in range(4)]
    data = pil_bytes(frames[0], "GIF", save_all=True, append_images=frames[1:], duration=120,
                     loop=0)
    r = gif.decode(data)
    assert len(r.frames) == 4 and r.kind == "animation" and r.frames[1].duration == 120
    ref = Image.open(io.BytesIO(data))
    for i, f in enumerate(r.frames):
        ref.seek(i)
        assert f.to("RGB").data == ref.convert("RGB").tobytes()


def test_jpeg_baseline_decode_close_to_pillow():
    for sub, mode in ((0, "RGB"), (2, "RGB"), (1, "RGB"), (0, "L")):
        im = gradient(mode, (61, 43))
        data = pil_bytes(im, "JPEG", quality=90, subsampling=sub) if mode == "RGB" \
            else pil_bytes(im, "JPEG", quality=90)
        same_pixels(jpeg.decode(data), Image.open(io.BytesIO(data)), tol=3)
    assert jpeg.jpeg_info(data) == (61, 43, 1)


def test_jpeg_restart_markers():
    im = gradient("RGB", (64, 48))
    data = pil_bytes(im, "JPEG", quality=85, restart_marker_blocks=2)
    same_pixels(jpeg.decode(data), Image.open(io.BytesIO(data)), tol=3)


def test_progressive_jpeg_is_unsupported():
    data = pil_bytes(gradient("RGB"), "JPEG", progressive=True)
    with pytest.raises(Unsupported):
        jpeg.decode(data)


def test_ico_round_trip():
    im = gradient("RGBA", (48, 48))
    data = ico.encode(Raster([Frame(48, 48, "RGBA", im.tobytes())]))
    back = Image.open(io.BytesIO(data))
    assert back.size == (48, 48)
    r = ico.decode(data)
    assert (r.first.width, r.first.mode) == (48, "RGBA")
    same_pixels(ico.decode(pil_bytes(im, "ICO", sizes=[(48, 48)])), im)


# ------------------------------------------------------------------ documents

MD = """# Title

Some **bold** and *italic* text with `code` and a [link](https://example.com).

- one
- two
  - nested

1. first
2. second

> quoted

```
print("hi")
```

| a | b |
|---|---|
| 1 | 2 |

---
"""


def test_markdown_to_html_and_back():
    blocks = docs.parse_markdown(MD)
    html = docs.to_html(blocks, "T")
    for frag in ("<h1>Title</h1>", "<strong>bold</strong>", "<em>italic</em>", "<code>code</code>",
                 '<a href="https://example.com">link</a>', "<ul>", "<ol>", "<blockquote>",
                 "<pre><code>print(&quot;hi&quot;)", "<table>", "<hr>"):
        assert frag.replace("&quot;", '"') in html.replace("&quot;", '"'), frag
    md2 = docs.to_markdown(docs.parse_html(html))
    assert "# Title" in md2 and "**bold**" in md2 and "[link](https://example.com)" in md2
    assert "- one" in md2 and "1. first" in md2


def test_html_to_text_skips_scripts():
    text = docs.to_text(docs.parse_html("<html><script>var x=1</script><h1>Hi</h1><p>Body "
                                        "<b>bold</b></p><ul><li>a</li></ul></html>"))
    assert "var x" not in text and "Hi" in text and "Body bold" in text and "• a" in text


def test_document_to_pdf_readable(tmp_path):
    import pymupdf
    p = tmp_path / "notes.md"
    p.write_text(MD * 20)       # long enough for several pages
    res = engine.convert_inplace(p, "pdf", Context(cfg=dict(CFG, prefer_native=True)))
    doc = pymupdf.open(res.outputs[0])
    assert doc.page_count >= 2
    assert "Title" in doc[0].get_text() and "bold" in doc[0].get_text()


def test_jpg_to_pdf_passthrough():
    import pymupdf
    data = pil_bytes(gradient("RGB", (80, 60)), "JPEG")
    from shapeshifter.native import pdfwrite
    pdf = pdfwrite.images_to_pdf(jpeg=data, jpeg_size=jpeg.jpeg_info(data))
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    assert doc.page_count == 1
    assert doc.extract_image(doc[0].get_images()[0][0])["image"] == data   # bytes untouched


# ---------------------------------------------------------------------- audio

def _wav(bits: int, rate=22050, channels=2, n=500) -> bytes:
    b = io.BytesIO()
    with wave.open(b, "wb") as w:
        w.setnchannels(channels)
        w.setsampwidth(bits // 8)
        w.setframerate(rate)
        if bits == 8:
            w.writeframes(bytes((i * 3) % 256 for i in range(n * channels)))
        elif bits == 16:
            w.writeframes(struct.pack(f"<{n * channels}h", *((i * 97) % 65536 - 32768
                                                             for i in range(n * channels))))
        else:
            w.writeframes(bytes((i * 13) % 256 for i in range(n * channels * 3)))
    return b.getvalue()


@pytest.mark.parametrize("bits", [8, 16, 24])
def test_wav_aiff_round_trip(bits):
    original = _wav(bits)
    aiff = audio.write_aiff(audio.read_wav(original))
    pcm = audio.read_aiff(aiff)
    assert (pcm.rate, pcm.channels, pcm.bits) == (22050, 2, bits)
    back = audio.write_wav(pcm)
    with wave.open(io.BytesIO(back)) as a, wave.open(io.BytesIO(original)) as b:
        assert a.readframes(10 ** 6) == b.readframes(10 ** 6)


def test_aiff_readable_by_ffmpeg(tmp_path):
    from shapeshifter.backends.media import probe_duration, ffprobe_path
    if not ffprobe_path():
        pytest.skip("ffprobe not installed")
    p = tmp_path / "t.aiff"
    p.write_bytes(audio.write_aiff(audio.read_wav(_wav(16, rate=8000, channels=1, n=8000))))
    assert abs(probe_duration(p) - 1.0) < 0.01


# ----------------------------------------------------------------------- data

def test_native_xlsx_interop(tmp_path):
    import openpyxl
    p = tmp_path / "ours.xlsx"
    xlsx.write(p, {"People": [["name", "age", "ok"], ["ann", 31, True], ["bob", 4.5, None]],
                   "Other": [["x"], ["y & <z>"]]})
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["People", "Other"]
    assert [c.value for c in wb["People"][2]] == ["ann", 31, True]
    assert wb["Other"]["A2"].value == "y & <z>"
    q = tmp_path / "theirs.xlsx"
    wb2 = openpyxl.Workbook()
    wb2.active.append(["k", "v"])
    wb2.active.append(["a", 1])
    wb2.active.append(["b", 2.5])
    wb2.save(q)
    assert xlsx.read(q) == {"Sheet": [["k", "v"], ["a", 1], ["b", 2.5]]}


def test_xml_ini_toml_conversions(tmp_path):
    import json
    import tomllib
    x = tmp_path / "cfg.xml"
    x.write_text('<config version="2"><item>a</item><item>b</item><name>demo</name></config>')
    res = engine.convert_inplace(x, "json", Context(cfg=CFG))
    assert json.loads(res.outputs[0].read_text()) == {
        "config": {"@version": 2, "item": ["a", "b"], "name": "demo"}}
    i = tmp_path / "app.ini"
    i.write_text("[server]\nport = 8080\nhost = example.org\n")
    res = engine.convert_inplace(i, "toml", Context(cfg=CFG))
    assert tomllib.loads(res.outputs[0].read_text()) == {"server": {"port": 8080,
                                                                    "host": "example.org"}}
    res = engine.convert_inplace(res.outputs[0], "xml", Context(cfg=CFG))
    assert "<port>8080</port>" in res.outputs[0].read_text()


# ------------------------------------------------------------ planner/fallback

@pytest.fixture
def prefer_native(monkeypatch):
    real = config.load
    monkeypatch.setattr(config, "load", lambda: dict(real(), prefer_native=True))
    planner.reset()
    yield
    monkeypatch.setattr(config, "load", real)
    planner.reset()


def test_prefer_native_routes(prefer_native):
    assert planner.plan("png")["bmp"].path[0].backend.name == "native"
    assert planner.plan("png")["webp"].path[0].backend.name == "images"   # no native WEBP


def test_fallback_when_native_cannot(tmp_path, prefer_native):
    p = tmp_path / "prog.jpg"
    p.write_bytes(pil_bytes(gradient("RGB", (40, 30)), "JPEG", progressive=True))
    assert planner.plan("jpg")["png"].path[0].backend.name == "native"
    res = engine.convert_inplace(p, "png", Context(cfg=dict(CFG, prefer_native=True)))
    assert Image.open(res.outputs[0]).size == (40, 30)      # Pillow took over


def test_native_end_to_end(tmp_path, prefer_native):
    src = tmp_path / "pic.png"
    gradient("RGBA").save(src)
    for target in ("bmp", "tga", "qoi", "ppm", "ico", "pdf"):
        res = engine.convert_inplace(src, target, Context(cfg=dict(CFG, prefer_native=True)))
        assert res.outputs[0].stat().st_size > 0, target

"""Combining several selected files into one (the inverse of every frame / sheet / page)."""

import subprocess
import sys
import zipfile
from pathlib import Path

import openpyxl
import pymupdf
import pytest
from PIL import Image, ImageDraw

from shapeshifter import combine, engine
from shapeshifter.backends import media
from shapeshifter.core import Context
from shapeshifter.native import docread, slides, slides_write

sys.path.insert(0, str(Path(__file__).parent))
from test_ocr import CFG                                  # noqa: E402

needs_ffmpeg = pytest.mark.skipif(not media.ffmpeg_path(), reason="ffmpeg not available")


def do(paths, fmt):
    return engine.combine_inplace([str(p) for p in paths], f"combine:{fmt}", Context(cfg=dict(CFG, original="recycle")))


def pictures(tmp_path, n=4, size=(120, 80), names=None):
    paths = []
    for i in range(n):
        im = Image.new("RGB", size, (50 * i, 100, 200 - 40 * i))
        ImageDraw.Draw(im).text((10, 30), f"pic {i + 1}", fill="white")
        p = tmp_path / (names[i] if names else f"frame{i + 1}.png")
        im.save(p)
        paths.append(p)
    return paths


class TestWhatIsOffered:
    def test_nothing_for_a_single_file(self, tmp_path):
        assert combine.targets_for(pictures(tmp_path, 1)) == []

    def test_pictures(self, tmp_path):
        t = combine.targets_for(pictures(tmp_path))
        assert {"gif", "webp", "tiff", "pdf", "pptx"} <= set(t)
        assert ("mp4" in t) == bool(media.ffmpeg_path())

    def test_other_kinds(self, tmp_path):
        pdfs = [tmp_path / f"{n}.pdf" for n in "ab"]
        for p in pdfs:
            doc = pymupdf.open()
            doc.new_page()
            doc.save(p)
        assert combine.targets_for(pdfs) == ["pdf"]
        assert combine.targets_for(pdfs + pictures(tmp_path, 1)) == ["pdf"]          # PDFs and pictures together
        csvs = [tmp_path / "a.csv", tmp_path / "b.tsv"]
        csvs[0].write_text("a,b\n1,2\n")
        csvs[1].write_text("c\td\n3\t4\n")
        assert combine.targets_for(csvs) == ["xlsx", "ods"]
        docs_ = [tmp_path / "a.md", tmp_path / "b.txt"]
        for d in docs_:
            d.write_text("hello")
        assert combine.targets_for(docs_) == ["pdf", "docx", "md"]

    def test_mixed_unrelated_files_offer_nothing(self, tmp_path):
        (tmp_path / "a.csv").write_text("a,b\n1,2\n")
        assert combine.targets_for(pictures(tmp_path, 1) + [tmp_path / "a.csv"]) == []

    def test_the_wheel_lists_combine_first_for_several_files(self, tmp_path):
        opts = engine.visible_options([str(p) for p in pictures(tmp_path)], dict(CFG))
        assert opts[0].kind == "combine" and opts[0].label.startswith("ONE ")
        assert "Combines all the selected files" in opts[0].description and "kept" in opts[0].description
        single = engine.visible_options([str(pictures(tmp_path)[0])], dict(CFG))
        assert not any(o.kind == "combine" for o in single)


class TestNaming:
    def test_natural_order(self):
        names = ["frame10.png", "frame2.png", "frame1.png"]
        assert [p.name for p in combine.natural_sorted(names)] == ["frame1.png", "frame2.png", "frame10.png"]

    def test_output_names(self):
        assert combine.output_name([Path("holiday_01.jpg"), Path("holiday_02.jpg")], "pdf") == "holiday (combined).pdf"
        assert combine.output_name([Path("cat.jpg"), Path("zebra.jpg"), Path("owl.jpg")], "pdf") == "cat and 2 more.pdf"


class TestPictures:
    def test_animated_gif_in_name_order(self, tmp_path):
        ps = pictures(tmp_path, 10)
        res = do(ps, "gif")
        out = res.outputs[0]
        im = Image.open(out)
        assert out.name == "frame (combined).gif" and im.n_frames == 10 and im.size == (120, 80)
        assert im.info["duration"] == 500
        im.seek(1)                                          # frame2, not frame10, comes second
        assert im.convert("RGB").getpixel((100, 70)) == pytest.approx((50, 100, 160), abs=12)

    def test_pictures_of_different_sizes_are_fitted(self, tmp_path):
        a, b = pictures(tmp_path, 2)
        Image.new("RGB", (60, 160), "green").save(b)
        out = do([a, b], "gif").outputs[0]
        assert Image.open(out).size == (120, 80)

    def test_multipage_tiff_and_animated_webp(self, tmp_path):
        ps = pictures(tmp_path, 3)
        assert Image.open(do(ps, "tiff").outputs[0]).n_frames == 3
        assert Image.open(do(ps, "webp").outputs[0]).n_frames == 3

    def test_one_pdf_page_per_picture(self, tmp_path):
        ps = pictures(tmp_path, 3)
        Image.new("RGB", (200, 100), "red").save(tmp_path / "wide.jpg")
        out = do(ps + [tmp_path / "wide.jpg"], "pdf").outputs[0]
        with pymupdf.open(out) as doc:
            assert len(doc) == 4
            assert all(max(p.rect.width, p.rect.height) == pytest.approx(842, abs=1) for p in doc)
            assert doc[3].rect.width > doc[3].rect.height              # the wide one stays wide
            assert all(p.get_images() for p in doc)

    def test_pictures_become_slides(self, tmp_path):
        ps = pictures(tmp_path, 3)
        deck = slides.read_pptx(do(ps, "pptx").outputs[0])
        assert len(deck.slides) == 3 and all(s.shapes[0].kind == "picture" for s in deck.slides)

    @needs_ffmpeg
    def test_a_slideshow_video(self, tmp_path):
        ps = pictures(tmp_path, 3)
        out = do(ps, "mp4").outputs[0]
        info = media.probe(out)
        assert info.has_video and (info.width, info.height) == (1280, 720)
        assert 8 <= info.duration <= 10                                # 3 pictures x 3 seconds

    def test_originals_are_kept(self, tmp_path):
        ps = pictures(tmp_path, 3)
        do(ps, "gif")
        assert all(p.exists() for p in ps)


class TestPdfs:
    def test_merging_pdfs_and_pictures(self, tmp_path):
        pdfs = []
        for n, pages in (("a", 2), ("b", 1)):
            doc = pymupdf.open()
            for i in range(pages):
                doc.new_page().insert_text((72, 100), f"{n} page {i + 1}", fontsize=20)
            doc.save(tmp_path / f"{n}.pdf")
            pdfs.append(tmp_path / f"{n}.pdf")
        pic = pictures(tmp_path, 1)[0]
        with pymupdf.open(do(pdfs + [pic], "pdf").outputs[0]) as doc:
            assert len(doc) == 4
            assert "a page 1" in doc[0].get_text() and "b page 1" in doc[2].get_text() and doc[3].get_images()


class TestTables:
    def test_a_workbook_with_a_sheet_per_file(self, tmp_path):
        (tmp_path / "sales.csv").write_text("item,qty\npen,3\n", encoding="utf-8")
        (tmp_path / "costs.tsv").write_text("name\tcost\nink\t10\n", encoding="utf-8")
        wb = openpyxl.Workbook()
        wb.active.title = "One"
        wb.active.append(["x"])
        wb.create_sheet("Two").append(["y"])
        wb.save(tmp_path / "multi.xlsx")
        out = do([tmp_path / n for n in ("sales.csv", "costs.tsv", "multi.xlsx")], "xlsx").outputs[0]
        book = openpyxl.load_workbook(out)
        assert book.sheetnames == ["costs", "multi - One", "multi - Two", "sales"]   # in file-name order
        assert [c.value for c in book["sales"][2]] == ["pen", 3]

    def test_ods_too(self, tmp_path):
        (tmp_path / "a.csv").write_text("p,q\n1,2\n", encoding="utf-8")
        (tmp_path / "b.csv").write_text("r\n3\n", encoding="utf-8")
        out = do([tmp_path / "a.csv", tmp_path / "b.csv"], "ods").outputs[0]
        from shapeshifter.native import tables as T
        assert [t.name for t in T.from_ods(out)] == ["a", "b"]


class TestDocuments:
    def make(self, tmp_path):
        (tmp_path / "1 intro.md").write_text("# Hello\n\nFirst **file**.\n", encoding="utf-8")
        (tmp_path / "2 notes.txt").write_text("Second file text.", encoding="utf-8")
        return [tmp_path / "1 intro.md", tmp_path / "2 notes.txt"]

    def test_markdown_docx_and_pdf(self, tmp_path):
        ps = self.make(tmp_path)
        md = do(ps, "md").outputs[0].read_text(encoding="utf-8")
        assert md.index("# 1 intro") < md.index("First **file**") < md.index("# 2 notes") < md.index("Second file text")
        blocks, _ = docread.docx_blocks(do(ps, "docx").outputs[0])
        assert [b.text for b in blocks if b.kind == "h1"][:2] == ["1 intro", "Hello"] or "1 intro" in [b.text for b in blocks]
        with pymupdf.open(do(ps, "pdf").outputs[0]) as doc:
            text = "".join(p.get_text() for p in doc)
            assert "First file" in text and "Second file text" in text

    def test_decks_are_merged_and_scaled(self, tmp_path):
        from shapeshifter.native import docs
        a = slides_write.deck_from_blocks([docs.Block("h2", "A1"), docs.Block("h2", "A2")], "A")
        b = slides_write.deck_from_blocks([docs.Block("h2", "B1")], "B")
        b.width, b.height = b.width / 2, b.height / 2
        for s in b.slides:
            for sh in s.shapes:
                sh.x, sh.y, sh.w, sh.h = sh.x / 2, sh.y / 2, sh.w / 2, sh.h / 2
        slides_write.write_pptx(a, tmp_path / "a.pptx")
        slides_write.write_pptx(b, tmp_path / "b.pptx")
        out = do([tmp_path / "a.pptx", tmp_path / "b.pptx"], "pptx").outputs[0]
        merged = slides.read_pptx(out)
        assert [s.title for s in merged.slides] == ["A1", "A2", "B1"]
        last = merged.slides[2].shapes[0]
        assert last.w == pytest.approx(a.slides[0].shapes[0].w, rel=0.01)
        assert slides_write.read_odp(do([tmp_path / "a.pptx", tmp_path / "b.pptx"], "odp").outputs[0]).slides[2].title == "B1"


@needs_ffmpeg
class TestMedia:
    def clip(self, path, seconds, audio=True, size="320x240", freq=440):
        cmd = [media.ffmpeg_path(), "-y", "-loglevel", "error", "-f", "lavfi", "-i", f"testsrc=duration={seconds}:size={size}:rate=15"]
        if audio:
            cmd += ["-f", "lavfi", "-i", f"sine=frequency={freq}:duration={seconds}", "-c:a", "aac"]
        subprocess.run(cmd + ["-c:v", "libx264", "-pix_fmt", "yuv420p", str(path)], check=True,
                       creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        return path

    def test_audio_clips_are_joined(self, tmp_path):
        wavs = []
        for i, f in enumerate((440, 660)):
            p = tmp_path / f"part{i + 1}.wav"
            subprocess.run([media.ffmpeg_path(), "-y", "-loglevel", "error", "-f", "lavfi", "-i", f"sine=frequency={f}:duration=2", str(p)],
                           check=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            wavs.append(p)
        out = do(wavs, "mp3").outputs[0]
        assert out.name == "part (combined).mp3"
        assert 3.8 <= media.probe(out).duration <= 4.4

    def test_video_clips_of_different_sizes_are_joined(self, tmp_path):
        a = self.clip(tmp_path / "a.mp4", 2)
        b = self.clip(tmp_path / "b.mp4", 3, size="160x120")
        info = media.probe(do([a, b], "mp4").outputs[0])
        assert (info.width, info.height) == (320, 240) and 4.6 <= info.duration <= 5.5 and info.has_audio

    def test_a_silent_clip_makes_the_result_silent(self, tmp_path):
        a = self.clip(tmp_path / "a.mp4", 2)
        b = self.clip(tmp_path / "b.mp4", 2, audio=False)
        info = media.probe(do([a, b], "mkv").outputs[0])
        assert info.has_video and not info.has_audio


class TestSafety:
    def test_risky_files_are_refused(self, tmp_path):
        a, b = pictures(tmp_path, 2)
        zipfile.ZipFile(tmp_path / "bomb.zip", "w").close()
        (tmp_path / "x.exe").write_bytes(b"MZ" + bytes(200))
        with pytest.raises(Exception):
            do([a, tmp_path / "x.exe"], "pdf")

    def test_an_impossible_combination_is_refused(self, tmp_path):
        with pytest.raises(Exception, match="can't be combined"):
            do(pictures(tmp_path, 2), "xlsx")


class TestWheel:
    def test_selecting_several_files_and_combining(self, tmp_path):
        import os
        import time
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        from PySide6.QtWidgets import QApplication
        from shapeshifter.wheel import WheelOverlay
        app = QApplication.instance() or QApplication([])
        ps = pictures(tmp_path, 3)
        w = WheelOverlay([str(p) for p in ps], dict(CFG, original="recycle", animation_ms=0))
        assert w.targets[0] == "combine:mp4"                       # no per-picture MP4, so it keeps its own entry
        assert not any(t.startswith("combine:") and t != "combine:mp4" for t in w.targets)
        pdf = w.option("pdf")                                      # "all in one PDF" hides inside the PDF entry
        assert [v.key for v in pdf.variants][-1] == "combine:pdf" and len(pdf.variants) >= 2
        w._commit("combine:pdf")
        end = time.time() + 30
        while time.time() < end and w.state != "done":
            app.processEvents()
            time.sleep(0.02)
        assert w.state == "done" and w.message == "frame (combined).pdf"
        assert all(p.exists() for p in ps)                         # never removed

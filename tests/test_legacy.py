"""The old binary Office formats (.doc, .xls, .ppt), read natively from the published specs."""

import sys
from pathlib import Path

import openpyxl
import pymupdf
import pytest

from shapeshifter import engine, planner
from shapeshifter.core import Context
from shapeshifter.native import legacy, ole, slides

sys.path.insert(0, str(Path(__file__).parent))
from legacy_fixtures import build_doc, build_ole, build_ppt, build_xls   # noqa: E402
from test_ocr import CFG                                                 # noqa: E402


def run(path, target):
    return engine.convert_inplace(path, target, Context(cfg=dict(CFG, original="keep"))).outputs


class TestContainer:
    def test_small_and_large_streams(self):
        small, large = b"hello " * 20, bytes(range(256)) * 40
        f = ole.OleFile(build_ole({"Small": small, "Large": large, "Another": b"x" * 100}))
        assert sorted(f.names()) == ["Another", "Large", "Small"]
        assert f.read("Small") == small and f.read("Large") == large and f.read("another") == b"x" * 100

    def test_not_ole(self):
        with pytest.raises(ole.OleError):
            ole.OleFile(b"not an ole file" * 100)


class TestWord:
    TEXT = "My Report\rFirst paragraph with café.\rName\x07Qty\x07\x07pen\x073\x07\x07After the table.\r"

    def test_text_title_paragraphs_and_table(self):
        blocks = legacy.doc_blocks(build_doc(self.TEXT))
        assert [(b.kind, b.text if b.kind != "table" else b.rows) for b in blocks] == [
            ("h1", "My Report"), ("p", "First paragraph with café."),
            ("table", [["Name", "Qty"], ["pen", "3"]]), ("p", "After the table.")]

    def test_unicode_pieces_and_large_documents(self):
        assert legacy.doc_blocks(build_doc("日本語のテキスト\r", unicode_text=True))[0].text == "日本語のテキスト"
        long = "".join(f"Paragraph {i} has some words in it.\r" for i in range(400))
        blocks = legacy.doc_blocks(build_doc(long))                  # bigger than the mini-stream cut-off
        assert len(blocks) == 400 and blocks[399].text == "Paragraph 399 has some words in it."

    def test_fields_keep_what_the_reader_sees(self):
        text = "See \x13 HYPERLINK \"x\" \x14the site\x15 now.\r"
        assert legacy.doc_blocks(build_doc(text))[0].text == "See the site now."

    def test_not_a_word_file(self):
        with pytest.raises(legacy.LegacyError):
            legacy.doc_text(build_xls({"S": [["a"]]}))

    def test_conversions(self, tmp_path):
        (tmp_path / "r.doc").write_bytes(build_doc(self.TEXT))
        md = run(tmp_path / "r.doc", "md")[0].read_text(encoding="utf-8")
        assert "# My Report" in md and "| pen | 3 |" in md
        (docx,) = run(tmp_path / "r.doc", "docx")
        from shapeshifter.native import docread
        assert [b.text for b in docread.docx_blocks(docx)[0] if b.kind == "p"][0] == "First paragraph with café."
        with pymupdf.open(run(tmp_path / "r.doc", "pdf")[0]) as doc:
            assert "First paragraph" in doc[0].get_text()
        assert any("Old binary" in w for w in planner.plan("doc")["docx"].warnings)


class TestExcel:
    def book(self):
        return {"Items": [["name", "qty", "price", "when", "ok"],
                          ["pen", 3, ("rk100", 125), ("date", 45000.0), True],
                          ["café ☕", ("rk", 7), 2.5, ("date", 45000.5), False],
                          [None, None, ("formula", 10.0), None, ("formula_str", "total")]],
                "Second": [["x"], ["y"]]}

    def test_cell_values(self):
        got = legacy.read_xls(build_xls(self.book()))
        assert list(got) == ["Items", "Second"]
        items = got["Items"]
        assert items[0] == ["name", "qty", "price", "when", "ok"]
        assert items[1][:3] == ["pen", 3, 1.25] and items[1][4] is True
        assert items[1][3] == "2023-03-15"                          # built-in date format
        assert items[2][0] == "café ☕" and items[2][1] == 7 and items[2][2] == 2.5
        assert items[2][3] == "2023-03-15 12:00:00"                 # a fraction of a day: date and time
        assert items[3][2] == 10 and items[3][4] == "total"          # formula results
        assert got["Second"] == [["x"], ["y"]]

    def test_a_long_string_table_continues_across_records(self):
        names = [f"a fairly long shared string number {i} with ünïcode" for i in range(60)]
        rows = [[n] for n in names]
        assert [r[0] for r in legacy.read_xls(build_xls({"S": rows}))["S"]] == names

    def test_not_excel(self):
        with pytest.raises(legacy.LegacyError):
            legacy.read_xls(build_doc("hi\r"))

    def test_conversions_both_ways(self, tmp_path):
        (tmp_path / "b.xls").write_bytes(build_xls(self.book()))
        (xlsx,) = run(tmp_path / "b.xls", "xlsx")
        wb = openpyxl.load_workbook(xlsx)
        assert wb.sheetnames == ["Items", "Second"] and [c.value for c in wb["Items"][2]][:3] == ["pen", 3, 1.25]
        csv_file = run(tmp_path / "b.xls", "csv")[0]
        assert csv_file.read_text(encoding="utf-8").splitlines()[1].startswith("pen,3,1.25,2023-03-15,True")
        assert len(run(tmp_path / "b.xls", "csv#sheets")) == 2
        assert "name" in run(tmp_path / "b.xls", "json")[0].read_text(encoding="utf-8")
        assert "| pen |" in run(tmp_path / "b.xls", "md")[0].read_text(encoding="utf-8") or "pen" in run(tmp_path / "b.xls", "md")[0].read_text(encoding="utf-8")
        assert run(tmp_path / "b.xls", "ods")[0].suffix == ".ods"
        assert any("Old binary" in w for w in planner.plan("xls")["xlsx"].warnings)

    def test_xlsm_is_read_like_xlsx(self, tmp_path):
        wb = openpyxl.Workbook()
        wb.active.append(["a", 1])
        wb.save(tmp_path / "m.xlsx")
        (tmp_path / "m.xlsm").write_bytes((tmp_path / "m.xlsx").read_bytes())
        assert run(tmp_path / "m.xlsm", "csv")[0].read_text(encoding="utf-8").strip() == "a,1"


class TestPowerPoint:
    DECK = [("Welcome", ["Intro line", "Second point"]), ("Plan", ["One", "Two", "Three"])]

    def test_slides_titles_and_text(self):
        got = legacy.read_ppt(build_ppt(self.DECK))
        assert [(s["title"], s["paragraphs"]) for s in got] == [(t, b) for t, b in self.DECK]

    def test_unicode_text(self):
        got = legacy.read_ppt(build_ppt([("Café", ["naïve ☕"])]))
        assert got[0]["paragraphs"] == ["naïve ☕"]

    def test_conversions_both_ways(self, tmp_path):
        (tmp_path / "d.ppt").write_bytes(build_ppt(self.DECK))
        (pptx,) = run(tmp_path / "d.ppt", "pptx")
        deck = slides.read_pptx(pptx)
        assert [s.title for s in deck.slides] == ["Welcome", "Plan"]
        md = run(tmp_path / "d.ppt", "md")[0].read_text(encoding="utf-8")
        assert "## Welcome" in md and "- Second point" in md
        with pymupdf.open(run(tmp_path / "d.ppt", "pdf")[0]) as doc:
            assert len(doc) == 2 and "Three" in doc[1].get_text()
        assert any("Old binary" in w for w in planner.plan("ppt")["pptx"].warnings)
        assert len(run(tmp_path / "d.ppt", "png")) == 2                  # and on into pictures


class TestRoutes:
    def test_longer_routes(self):
        for src, dst in (("doc", "epub"), ("doc", "rtf"), ("xls", "pdf"), ("xls", "docx"), ("ppt", "odp"), ("xls", "ods"),
                         ("doc", "csv"), ("ppt", "docx"), ("xls", "pptx")):
            assert planner.route(src, dst) is not None, (src, dst)

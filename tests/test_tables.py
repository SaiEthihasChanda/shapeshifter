"""Tables between documents and data, in both directions and through chained steps."""

import csv
import json
import sys
import zipfile
from pathlib import Path

import openpyxl
import pymupdf
import pytest
from PIL import Image, ImageDraw, ImageFont

from shapeshifter import engine, ocr, planner
from shapeshifter.core import Context
from shapeshifter.native import tables as T

sys.path.insert(0, str(Path(__file__).parent))
from test_ocr import CFG, needs_ocr                       # noqa: E402


def run(path, target, tmp_keep=True):
    cfg = dict(CFG, original="keep")
    return engine.convert_inplace(path, target, Context(cfg=cfg)).outputs


def make_calendar_pdf(path: Path, pages: int = 2, per_page: int = 12) -> Path:
    """A ruled table spread over pages, header repeated on each (like a printed timetable)."""
    header = ["Date", "Day", "B.Tech.", "B.Com."]
    doc = pymupdf.open()
    n = 0
    for _ in range(pages):
        pg = doc.new_page(width=595, height=842)
        rows = [header] + [[f"{14 + (n + i):02d}-07-2025", "Mon", str(n + i), str(n + i + 1)] for i in range(per_page)]
        n += per_page
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                rect = pymupdf.Rect(50 + c * 120, 60 + r * 24, 50 + (c + 1) * 120, 60 + (r + 1) * 24)
                pg.draw_rect(rect, width=0.8)
                pg.insert_text((rect.x0 + 6, rect.y1 - 7), v, fontsize=10)
    doc.save(path)
    doc.close()
    return path


def make_xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["item", "qty", "note"])
    ws.append(["pencil", 12, "00123"])
    ws.append(["pen", 3, "a, b"])
    wb.create_sheet("Costs").append(["x", "y"])
    wb["Costs"].append([1.5, 2])
    wb.save(path)
    return path


class TestFromPdf:
    def test_cells_keep_their_punctuation_and_a_repeated_header_joins_the_pages(self, tmp_path):
        pdf = make_calendar_pdf(tmp_path / "cal.pdf")
        (xlsx,) = run(pdf, "xlsx")
        ws = openpyxl.load_workbook(xlsx).active
        assert openpyxl.load_workbook(xlsx).sheetnames == ["Pages 1-2"]
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        assert rows[0] == ["Date", "Day", "B.Tech.", "B.Com."]
        assert rows[1][0] == "14-07-2025" and rows[1][2:] == [0, 1]          # numbers become numbers
        assert len(rows) == 25 and rows[-1][0] == "37-07-2025"               # both pages, one header

    def test_csv_tsv_and_json(self, tmp_path):
        pdf = make_calendar_pdf(tmp_path / "cal.pdf", pages=1)
        (c,) = run(pdf, "csv")
        assert list(csv.reader(open(c, encoding="utf-8")))[1][:2] == ["14-07-2025", "Mon"]
        (t,) = run(pdf, "tsv")
        assert open(t, encoding="utf-8").read().splitlines()[0] == "Date\tDay\tB.Tech.\tB.Com."
        (j,) = run(pdf, "json")
        data = json.load(open(j, encoding="utf-8"))
        assert data[0] == {"Date": "14-07-2025", "Day": "Mon", "B.Tech.": 0, "B.Com.": 1}

    def test_a_pdf_without_tables_says_so(self, tmp_path):
        doc = pymupdf.open()
        doc.new_page().insert_text((72, 100), "Just a sentence, no table here.", fontsize=12)
        doc.save(tmp_path / "plain.pdf")
        with pytest.raises(Exception, match="No tables"):
            run(tmp_path / "plain.pdf", "xlsx")

    def test_pdf_tables_carry_a_layout_warning(self):
        opt = planner.plan("pdf")["xlsx"]
        assert any("page layout" in w for w in opt.warnings)
        assert opt.path[0].backend.name == "tables"


class TestOtherSources:
    def md_table(self):
        return "# Prices\n\n| item | qty |\n|---|---|\n| pen | 3 |\n| ink | 10 |\n\n## Costs\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"

    def test_markdown_and_html_tables_with_their_headings(self, tmp_path):
        (tmp_path / "t.md").write_text(self.md_table(), encoding="utf-8")
        (xlsx,) = run(tmp_path / "t.md", "xlsx")
        wb = openpyxl.load_workbook(xlsx)
        assert wb.sheetnames == ["Prices", "Costs"]
        assert [c.value for c in wb["Prices"][2]] == ["pen", 3]
        (html,) = run(tmp_path / "t.md", "html")
        outs = run(html, "csv")                      # HTML -> CSV: one file per table
        assert sorted(o.name for o in outs) == ["t - Costs.csv", "t - Prices.csv"]

    def test_word_tables(self, tmp_path):
        (tmp_path / "t.md").write_text(self.md_table(), encoding="utf-8")
        (docx,) = run(tmp_path / "t.md", "docx")
        tabs = T.from_docx(docx)
        assert [t.name for t in tabs] == ["Prices", "Costs"] and tabs[0].rows[1] == ["pen", "3"]
        outs = run(docx, "json")
        assert json.load(open(outs[0], encoding="utf-8"))["Prices"][0] == {"item": "pen", "qty": 3}

    def test_opendocument_text_tables(self, tmp_path):
        content = ('<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                   'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
                   'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
                   '<table:table table:name="T1"><table:table-row><table:table-cell><text:p>a</text:p></table:table-cell>'
                   '<table:table-cell><text:p>b</text:p></table:table-cell></table:table-row>'
                   '<table:table-row><table:table-cell><text:p>1</text:p></table:table-cell>'
                   '<table:table-cell><text:p>2</text:p></table:table-cell></table:table-row></table:table>'
                   '</office:text></office:body></office:document-content>')
        path = tmp_path / "t.odt"
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("mimetype", "application/vnd.oasis.opendocument.text")
            z.writestr("content.xml", content)
        (tab,) = T.from_odt(path)
        assert tab.rows == [["a", "b"], ["1", "2"]]


class TestSpreadsheetsToDocuments:
    """The way back: data -> pages."""

    def test_to_html_markdown_word_and_pdf(self, tmp_path):
        x = make_xlsx(tmp_path / "book.xlsx")
        (html,) = run(x, "html")
        text = html.read_text(encoding="utf-8")
        assert "<table>" in text and "pencil" in text and "Costs" in text and "00123" in text
        (md,) = run(x, "md")
        assert "| item | qty | note |" in md.read_text(encoding="utf-8")
        (docx,) = run(x, "docx")
        assert [t.rows[0] for t in T.from_docx(docx)] == [["item", "qty", "note"], ["x", "y"]]
        (pdf,) = run(x, "pdf")
        with pymupdf.open(pdf) as doc:
            assert "pencil" in doc[0].get_text() and "Costs" in doc[0].get_text()

    def test_csv_to_a_pdf_table(self, tmp_path):
        (tmp_path / "d.csv").write_text("name,score\nAda,10\nGrace,9\n", encoding="utf-8")
        (pdf,) = run(tmp_path / "d.csv", "pdf")
        with pymupdf.open(pdf) as doc:
            text = doc[0].get_text()
        assert "Ada" in text and "Grace" in text


class TestOds:
    def test_round_trip_keeps_sheets_numbers_and_text(self, tmp_path):
        x = make_xlsx(tmp_path / "book.xlsx")
        (ods,) = run(x, "ods")
        with zipfile.ZipFile(ods) as z:
            assert z.namelist()[0] == "mimetype"                      # must come first, uncompressed
            assert z.read("mimetype") == b"application/vnd.oasis.opendocument.spreadsheet"
        tabs = T.from_ods(ods)
        assert [t.name for t in tabs] == ["Sales", "Costs"]
        assert tabs[0].rows[1] == ["pencil", 12, "00123"]             # leading zeros stay text
        (back,) = run(ods, "xlsx")
        assert [c.value for c in openpyxl.load_workbook(back)["Sales"][3]] == ["pen", 3, "a, b"]
        outs = run(ods, "csv")
        assert outs[0].read_text(encoding="utf-8").splitlines()[1] == "pencil,12,00123"


class TestChainedRoutes:
    def test_longer_routes_exist(self):
        # a spreadsheet has no direct EPUB/ODT writer: it goes through a document step
        for src, dst in (("csv", "odt"), ("xlsx", "epub"), ("sqlite", "pdf"), ("pdf", "ods"),
                         ("docx", "csv"), ("md", "xlsx")):
            assert planner.route(src, dst) is not None, (src, dst)

    def test_html_is_not_used_as_a_stepping_stone_for_tables(self):
        # finding tables only makes sense on the user's own document, never on an intermediate
        opt = planner.route("txt", "csv")
        assert opt is None or all(e.backend.name != "tables" for e in opt.path[1:])

    def test_sqlite_to_a_word_table_through_a_spreadsheet(self, tmp_path):
        import sqlite3
        db = tmp_path / "shop.sqlite"
        con = sqlite3.connect(db)
        con.execute("create table items(name text, qty integer)")
        con.executemany("insert into items values(?,?)", [("pen", 3), ("ink", 10)])
        con.commit()
        con.close()
        (docx,) = run(db, "docx")
        assert T.from_docx(docx)[0].rows == [["name", "qty"], ["pen", "3"], ["ink", "10"]]


class TestFromPictures:
    def test_words_are_grouped_into_rows_and_columns(self):
        words = []
        rows = [["Item", "Qty", "Price"], ["Pencil", "12", "1.50"], ["Notebook", "5", "3.25"]]
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                words.append((text, 40 + [0, 380, 620][c], 30 + r * 70, 28 * len(text) * 0.55, 28))
        assert T.rows_from_words(words) == rows

    def test_prose_is_not_a_table(self):
        words = [(w, 40 + i * 70, 30, 60, 24) for i, w in enumerate("just one line of text".split())]
        assert T.rows_from_words(words) == []

    @needs_ocr
    def test_a_picture_of_a_table_becomes_a_spreadsheet(self, tmp_path):
        from test_ocr import ARIAL
        f = ImageFont.truetype(str(ARIAL), 28)
        im = Image.new("RGB", (900, 420), "white")
        d = ImageDraw.Draw(im)
        rows = [["Item", "Qty", "Price"], ["Pencil", "12", "1.50"], ["Notebook", "5", "3.25"], ["Ruler", "8", "2.00"]]
        for r, row in enumerate(rows):
            for c, v in enumerate(row):
                d.text((40 + [0, 380, 620][c], 30 + r * 70), v, font=f, fill="black")
        im.save(tmp_path / "t.png")
        (xlsx,) = run(tmp_path / "t.png", "xlsx")
        got = [[c.value for c in r] for r in openpyxl.load_workbook(xlsx).active.iter_rows()]
        assert len(got) == 4 and len(got[0]) == 3
        assert got[1][:2] == ["Pencil", 12] and got[3][2] == 2.0
        opt = planner.plan("png")["xlsx"]
        assert opt.category == "inferred" and any("OCR" in w for w in opt.warnings)


class TestCoerce:
    def test_only_safe_numbers_are_converted(self):
        assert T.coerce("12") == 12 and T.coerce("3.50") == 3.5 and T.coerce("-4") == -4
        for keep in ("0123", "1,234", "+91 98765", "12-03-2004", "1e5", "", "abc"):
            assert T.coerce(keep) == keep

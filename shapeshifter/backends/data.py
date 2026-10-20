"""Structured data conversions: CSV, TSV, XLSX, SQLite, JSON, YAML, TOML, XML, INI.

XLSX uses openpyxl when installed, otherwise (or with "prefer native") the built-in
reader/writer in native/xlsx.py.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from .base import Backend, Context, ConversionError, target_ext

TABLE = {"csv", "tsv", "xlsx", "sqlite"}
TREE = {"json", "yaml", "toml", "xml", "ini"}
WRITABLE = {"csv", "tsv", "json", "yaml", "xlsx", "toml", "xml"}

_NUM = re.compile(r"-?(0|[1-9]\d*)(\.\d+)?([eE][-+]?\d+)?")

_native_xlsx = False   # set per conversion from the "prefer_native" setting


class DataBackend(Backend):
    name = "data"

    def __init__(self):
        from importlib.util import find_spec
        self._yaml = find_spec("yaml") is not None

    def _ok(self, ext):
        return ext != "yaml" or self._yaml

    def sources(self):
        return {e for e in TABLE | TREE if self._ok(e)}

    def targets(self, src):
        return {e for e in WRITABLE if self._ok(e)} - {src}

    def convert(self, src: Path, dst: Path, ctx: Context) -> list[Path]:
        global _native_xlsx
        from importlib.util import find_spec
        _native_xlsx = bool(ctx.cfg.get("prefer_native")) or find_spec("openpyxl") is None
        s = ctx.src_fmt or target_ext(src)
        t = ctx.dst_fmt or target_ext(dst)
        try:
            if s in TABLE:
                sheets = _read_table(src, s)
                if ctx.variant == "sheets" and t in ("csv", "tsv") and len(sheets) > 1:
                    outs = []
                    for name, rows in sheets.items():
                        safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name)).strip() or "Sheet"
                        out = dst.with_name(f"{dst.stem} - {safe}{dst.suffix}")
                        _write_table(out, t, {name: rows})
                        outs.append(out)
                    ctx.progress(1.0)
                    return outs
                if t in TABLE:
                    _write_table(dst, t, sheets)
                else:
                    tree = {name: _records(rows) for name, rows in sheets.items()}
                    _write_tree(dst, t, next(iter(tree.values())) if len(tree) == 1 else tree)
            else:
                tree = _read_tree(src, s)
                if t in TABLE:
                    _write_table(dst, t, {"Sheet1": _to_rows(tree)})
                else:
                    _write_tree(dst, t, tree)
        except ConversionError:
            raise
        except Exception as e:
            raise ConversionError(f"Couldn't convert data: {e}") from e
        ctx.progress(1.0)
        return [dst]


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _read_table(path: Path, ext: str) -> dict[str, list[list]]:
    if ext == "sqlite":
        import sqlite3
        # Read-only and immutable: never writes to (or journals next to) the database.
        con = sqlite3.connect(f"{path.as_uri()}?mode=ro&immutable=1", uri=True)
        try:
            names = [r[0] for r in con.execute(
                "SELECT name FROM sqlite_master WHERE type IN ('table','view') "
                "AND name NOT LIKE 'sqlite_%' ORDER BY type, name")]
            sheets = {}
            for name in names:
                cur = con.execute(f'SELECT * FROM "{name.replace(chr(34), chr(34) * 2)}"')
                header = [d[0] for d in cur.description]
                rows = [[v.hex() if isinstance(v, bytes) else v for v in r] for r in cur]
                sheets[name] = [header] + rows
            if not sheets:
                raise ConversionError("The database has no tables.")
            return sheets
        finally:
            con.close()
    if ext == "xlsx" and _native_xlsx:
        from ..native import xlsx
        return xlsx.read(path)
    if ext == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = {}
            for ws in wb.worksheets:
                rows = [list(r) for r in ws.iter_rows(values_only=True)]
                while rows and all(v is None for v in rows[-1]):
                    rows.pop()
                sheets[ws.title] = rows
            return sheets
        finally:
            wb.close()
    text = _read_text(path)
    if ext == "tsv":
        delim = "\t"
    else:
        try:
            delim = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delim = ","
    rows = [[_infer(v) for v in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]
    return {"Sheet1": rows}


def _infer(v: str):
    if _NUM.fullmatch(v):
        return float(v) if any(c in v for c in ".eE") else int(v)
    return v


def _records(rows: list[list]) -> list[dict]:
    if not rows:
        return []
    header = [str(h) if h is not None else f"column{i + 1}" for i, h in enumerate(rows[0])]
    out = []
    for row in rows[1:]:
        row = list(row) + [None] * (len(header) - len(row))
        out.append({h: _jsonable(v) for h, v in zip(header, row)})
    return out


def _jsonable(v):
    if v is None or isinstance(v, (str, int, float, bool)):
        return v
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def _read_tree(path: Path, ext: str):
    text = _read_text(path)
    if ext == "json":
        return json.loads(text)
    if ext == "yaml":
        import yaml
        docs = list(yaml.safe_load_all(text))
        return docs[0] if len(docs) == 1 else docs
    if ext == "toml":
        import tomllib
        return tomllib.loads(text)
    if ext == "xml":
        from xml.etree import ElementTree as ET
        root = ET.fromstring(text.encode("utf-8"))
        return {_xml_tag(root.tag): _xml_to_tree(root)}
    if ext == "ini":
        import configparser
        cp = configparser.ConfigParser(interpolation=None)
        cp.optionxform = str
        cp.read_string(text if re.match(r"\s*\[", text) else "[DEFAULT]\n" + text)
        out = {s: {k: _infer(v) for k, v in cp[s].items()} for s in cp.sections()}
        if cp.defaults():
            out = {"DEFAULT": {k: _infer(v) for k, v in cp.defaults().items()}, **out}
        return out
    raise ConversionError(f"Unsupported source {ext}")


def _xml_tag(tag: str) -> str:
    return tag.split("}", 1)[-1]


def _xml_to_tree(el):
    """<a x="1"><b>t</b><b>u</b></a> -> {"@x": 1, "b": ["t", "u"]}."""
    node: dict = {f"@{_xml_tag(k)}": _infer(v) for k, v in el.attrib.items()}
    for child in el:
        key = _xml_tag(child.tag)
        val = _xml_to_tree(child)
        if key in node:
            if not isinstance(node[key], list):
                node[key] = [node[key]]
            node[key].append(val)
        else:
            node[key] = val
    text = (el.text or "").strip()
    if text:
        if not node:
            return _infer(text)
        node["#text"] = text
    return node or None


def _tree_to_xml(tag: str, value, indent: int = 0) -> str:
    from xml.sax.saxutils import escape, quoteattr
    pad = "  " * indent
    tag = re.sub(r"[^\w.-]", "_", str(tag)) or "item"
    if tag[0].isdigit() or tag[0] in ".-":
        tag = "_" + tag
    if isinstance(value, list):
        return "".join(_tree_to_xml(tag, v, indent) for v in value)
    if isinstance(value, dict):
        attrs = "".join(f" {k[1:]}={quoteattr(str(v))}" for k, v in value.items()
                        if str(k).startswith("@"))
        kids = {k: v for k, v in value.items() if not str(k).startswith("@") and k != "#text"}
        text = escape(str(value["#text"])) if "#text" in value else ""
        if not kids:
            return f"{pad}<{tag}{attrs}>{text}</{tag}>\n"
        inner = "".join(_tree_to_xml(k, v, indent + 1) for k, v in kids.items())
        return f"{pad}<{tag}{attrs}>{text}\n{inner}{pad}</{tag}>\n"
    if value is None:
        return f"{pad}<{tag}/>\n"
    return f"{pad}<{tag}>{escape(str(value).lower() if isinstance(value, bool) else str(value))}</{tag}>\n"


def _toml_value(v) -> str:
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(v)
    if v is None:
        return '""'
    if isinstance(v, list):
        return "[" + ", ".join(_toml_value(x) for x in v) + "]"
    if isinstance(v, dict):
        return "{ " + ", ".join(f"{_toml_key(k)} = {_toml_value(x)}" for k, x in v.items()) + " }"
    return json.dumps(str(v), ensure_ascii=False)


def _toml_key(k) -> str:
    k = str(k)
    return k if re.fullmatch(r"[A-Za-z0-9_-]+", k) else json.dumps(k, ensure_ascii=False)


def _dump_toml(tree, prefix: str = "") -> str:
    if not isinstance(tree, dict):
        if isinstance(tree, list) and all(isinstance(x, dict) for x in tree):
            tree = {"items": tree}
        else:
            tree = {"value": tree}
    lines, tables = [], []
    for k, v in tree.items():
        if isinstance(v, dict):
            tables.append((k, v))
        elif isinstance(v, list) and v and all(isinstance(x, dict) for x in v):
            tables.append((k, v))
        else:
            lines.append(f"{_toml_key(k)} = {_toml_value(v)}")
    out = "\n".join(lines) + ("\n" if lines else "")
    for k, v in tables:
        name = f"{prefix}.{_toml_key(k)}" if prefix else _toml_key(k)
        if isinstance(v, list):
            for item in v:
                out += f"\n[[{name}]]\n" + _dump_toml(item, name)
        else:
            out += f"\n[{name}]\n" + _dump_toml(v, name)
    return out


def _to_rows(tree) -> list[list]:
    """Flatten a JSON-like structure into a table."""
    def cell(v):
        return json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v

    if isinstance(tree, dict):
        if tree and all(isinstance(v, dict) for v in tree.values()):
            tree = [{"key": k, **v} for k, v in tree.items()]
        elif tree and all(isinstance(v, list) for v in tree.values()) and len(tree) == 1:
            tree = next(iter(tree.values()))
        else:
            tree = [tree]
    if not isinstance(tree, list):
        return [["value"], [cell(tree)]]
    if all(isinstance(r, dict) for r in tree):
        header: list[str] = []
        for r in tree:
            for k in r:
                if k not in header:
                    header.append(k)
        return [header] + [[cell(r.get(h)) for h in header] for r in tree]
    if all(isinstance(r, list) for r in tree):
        return [[cell(v) for v in r] for r in tree]
    return [["value"]] + [[cell(v)] for v in tree]


def _write_table(path: Path, ext: str, sheets: dict[str, list[list]]):
    if ext == "xlsx" and _native_xlsx:
        from ..native import xlsx
        xlsx.write(path, {n: [[_jsonable(v) for v in r] for r in rows] for n, rows in sheets.items()})
        return
    if ext == "xlsx":
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name, rows in sheets.items():
            ws = wb.create_sheet(title=str(name)[:31] or "Sheet1")
            for row in rows:
                ws.append([_jsonable(v) if not isinstance(v, (int, float, str)) else v for v in row])
        wb.save(path)
        return
    rows = next(iter(sheets.values()), [])
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter="\t" if ext == "tsv" else ",")
        for row in rows:
            w.writerow(["" if v is None else _jsonable(v) for v in row])


def _write_tree(path: Path, ext: str, tree):
    if ext == "json":
        path.write_text(json.dumps(tree, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    elif ext == "yaml":
        import yaml
        path.write_text(yaml.safe_dump(json.loads(json.dumps(tree, default=str)), sort_keys=False,
                                       allow_unicode=True), encoding="utf-8")
    elif ext == "toml":
        path.write_text(_dump_toml(json.loads(json.dumps(tree, default=str))), encoding="utf-8")
    elif ext == "xml":
        tree = json.loads(json.dumps(tree, default=str))
        if isinstance(tree, dict) and len(tree) == 1:
            (root, body), = tree.items()
        else:
            root, body = "root", ({"item": tree} if isinstance(tree, list) else tree)
        path.write_text('<?xml version="1.0" encoding="UTF-8"?>\n' + _tree_to_xml(root, body),
                        encoding="utf-8")
    else:
        raise ConversionError(f"Unsupported target {ext}")
